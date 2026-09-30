import base64
import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from fpdf import FPDF

from .sanitize import sanitize_text


def safe_filename(
    document_type: str,
) -> str:

    name = re.sub(
        r"[^A-Za-z0-9]+",
        "_",
        document_type.strip(),
    )

    name = name.strip("_")

    if not name:
        name = "legal_document"

    return name[:80]


def format_txt(
    text: str,
) -> bytes:

    clean_text = sanitize_text(text)

    return clean_text.encode(
        "utf-8"
    )


def configure_docx_font(
    document: Document,
):

    styles = document.styles

    normal_style = styles["Normal"]

    normal_style.font.name = (
        "Times New Roman"
    )

    normal_style.font.size = Pt(11)

    normal_style._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Times New Roman",
    )


def add_logo_to_docx(
    document: Document,
    logo_base64: str | None,
):

    if not logo_base64:
        return

    try:

        encoded_data = logo_base64.split(
            ",",
            1,
        )[-1]

        image_data = base64.b64decode(
            encoded_data
        )

        image_stream = io.BytesIO(
            image_data
        )

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.add_run().add_picture(
            image_stream,
            width=Inches(1.5),
        )

    except Exception:
        # Invalid logo should not prevent
        # the document from being generated.
        return


def format_docx(
    text: str,
    doc_type: str,
    terms: str = "",
    logo_base64: str | None = None,
) -> bytes:

    document = Document()

    configure_docx_font(
        document
    )

    # Logo
    add_logo_to_docx(
        document,
        logo_base64,
    )

    # Title
    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title.add_run(
        sanitize_text(
            doc_type
        ).upper()
    )

    title_run.bold = True
    title_run.font.name = (
        "Times New Roman"
    )
    title_run.font.size = Pt(16)

    # Document body
    clean_text = sanitize_text(text)

    lines = clean_text.split("\n")

    for raw_line in lines:

        line = raw_line.strip()

        if not line:
            document.add_paragraph()
            continue

        paragraph = document.add_paragraph()

        is_numbered_section = bool(
            re.match(
                r"^(SECTION\s+)?\d+[.)]\s+",
                line,
                re.IGNORECASE,
            )
        )

        is_heading = (
            line.isupper()
            and len(line) < 100
        )

        if (
            is_numbered_section
            or is_heading
        ):

            run = paragraph.add_run(
                line
            )

            run.bold = True

        else:

            paragraph.add_run(
                line
            )

        paragraph.paragraph_format.space_after = (
            Pt(6)
        )

    # Terms table
    clean_terms = [
        item.strip()
        for item in terms.split(";")
        if item.strip()
    ]

    if clean_terms:

        document.add_heading(
            "Key Terms",
            level=2,
        )

        table = document.add_table(
            rows=1,
            cols=2,
        )

        table.style = "Table Grid"

        table.alignment = (
            WD_TABLE_ALIGNMENT.CENTER
        )

        header_cells = table.rows[0].cells

        header_cells[0].text = "#"
        header_cells[1].text = "Term"

        for index, term in enumerate(
            clean_terms,
            start=1,
        ):

            cells = table.add_row().cells

            cells[0].text = str(index)

            cells[1].text = sanitize_text(
                term
            )

    # Footer
    section = document.sections[0]

    footer = (
        section.footer.paragraphs[0]
    )

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer.add_run(
        "LegalEase - AI-generated draft | "
        "Review with a qualified legal professional "
        "before signing."
    )

    output = io.BytesIO()

    document.save(output)

    return output.getvalue()


class LegalEasePDF(FPDF):

    def header(self):

        self.set_font(
            "Helvetica",
            "B",
            12,
        )

        self.cell(
            0,
            8,
            "LegalEase",
            align="C",
        )

        self.ln(8)

        self.set_draw_color(
            160,
            160,
            160,
        )

        self.line(
            15,
            self.get_y(),
            195,
            self.get_y(),
        )

        self.ln(5)

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            "I",
            8,
        )

        self.cell(
            0,
            8,
            "LegalEase - AI-generated draft; "
            "professional legal review recommended.",
            align="C",
        )


def format_pdf(
    text: str,
    doc_type: str,
) -> bytes:

    pdf = LegalEasePDF()

    pdf.set_title(
        doc_type
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=20,
    )

    pdf.add_page()

    # Title
    pdf.set_font(
        "Helvetica",
        "B",
        16,
    )

    pdf.multi_cell(
        0,
        9,
        sanitize_text(
            doc_type
        ).upper(),
        align="C",
    )

    pdf.ln(4)

    # Body
    clean_text = sanitize_text(text)

    for raw_line in clean_text.split(
        "\n"
    ):

        line = raw_line.strip()

        if not line:

            pdf.ln(3)

            continue

        is_numbered = bool(
            re.match(
                r"^(SECTION\s+)?\d+[.)]\s+",
                line,
                re.IGNORECASE,
            )
        )

        is_heading = (
            line.isupper()
            and len(line) < 100
        )

        if (
            is_numbered
            or is_heading
        ):

            pdf.set_font(
                "Helvetica",
                "B",
                11,
            )

        else:

            pdf.set_font(
                "Helvetica",
                "",
                11,
            )

        pdf.multi_cell(
            0,
            6,
            line,
        )

        pdf.ln(1)

    return bytes(
        pdf.output()
    )


def export_document(
    fmt: str,
    text: str,
    doc_type: str,
    terms: str = "",
    logo_base64: str | None = None,
):

    fmt = fmt.lower().strip()

    filename = safe_filename(
        doc_type
    )

    if fmt == "txt":

        return (
            format_txt(text),
            "text/plain; charset=utf-8",
            f"{filename}.txt",
        )

    if fmt == "docx":

        return (
            format_docx(
                text,
                doc_type,
                terms,
                logo_base64,
            ),
            (
                "application/"
                "vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            f"{filename}.docx",
        )

    if fmt == "pdf":

        return (
            format_pdf(
                text,
                doc_type,
            ),
            "application/pdf",
            f"{filename}.pdf",
        )

    raise ValueError(
        "Unsupported format. "
        "Use txt, docx, or pdf."
    )